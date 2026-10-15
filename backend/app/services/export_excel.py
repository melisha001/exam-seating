"""
services/export_excel.py — Generate an Excel seating plan using openpyxl.

One sheet per room.  Each cell = one physical seat.
Cells are colour-coded by department with roll number + department code.
A summary sheet lists rooms, utilisation, and warnings.
"""

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

from app.models.seating_plan import SeatingPlan, SeatAssignment
from app.models.room import Room


# Hex fills per department (no leading #, openpyxl uses AARRGGBB or RRGGBB)
_DEPT_FILLS = [
    "DBEAFE", "DCFCE7", "FEF9C3", "FCE7F3", "EDE9FE",
    "FFEDD5", "CFFAFE", "F0FDF4", "FEF2F2", "F5F3FF",
]
_EMPTY_FILL = "F3F4F6"
_HEADER_FILL = "1E3A5F"   # dark blue
_HEADER_FONT_COLOR = "FFFFFF"

_thin = Side(style="thin", color="AAAAAA")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _make_fill(hex_colour: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_colour)


def _dept_fill(dept_id: int | None) -> PatternFill:
    if dept_id is None:
        return _make_fill(_EMPTY_FILL)
    return _make_fill(_DEPT_FILLS[(dept_id - 1) % len(_DEPT_FILLS)])


def generate_excel(
    plan: SeatingPlan,
    rooms: list[Room],
    seats_by_room: dict[int, list[SeatAssignment]],
) -> bytes:
    """
    Build an Excel workbook and return it as bytes.
    """
    wb = Workbook()
    wb.remove(wb.active)   # remove default blank sheet

    # ── Per-room sheets ───────────────────────────────────────────────────────
    for room in rooms:
        room_seats = seats_by_room.get(room.id, [])
        if not room_seats:
            continue

        ws = wb.create_sheet(title=room.name[:31])   # sheet name ≤31 chars
        max_spb = max(
            (room.get_bench_seats(r, c) for r in range(room.rows) for c in range(room.columns)),
            default=room.seats_per_bench,
        )

        # Header rows
        invig1_str = room.invigilator_name or "—"
        if room.invigilator1_subject:
            invig1_str += f" ({room.invigilator1_subject})"
        invig2_str = room.invigilator2_name or "—"
        if room.invigilator2_subject:
            invig2_str += f" ({room.invigilator2_subject})"

        used = sum(1 for s in room_seats if not s.is_empty)
        cap = room.capacity
        util = f"{round(used / cap * 100, 1)}%" if cap else "—"
        breakdown_str = f" ({room.seater_breakdown})" if getattr(room, "seater_breakdown", None) else ""

        total_excel_cols = max(1, room.columns * max_spb)
        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1, end_column=total_excel_cols)
        h1 = ws.cell(1, 1,
                     f"Hall: {room.name} — {cap} seats{breakdown_str}  |  In-Charge 1: {invig1_str}  |  In-Charge 2: {invig2_str}  |  "
                     f"Seated: {used}/{cap}  |  Utilisation: {util}")
        h1.font = Font(bold=True, color=_HEADER_FONT_COLOR, size=11)
        h1.fill = _make_fill(_HEADER_FILL)
        h1.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        # Column width
        col_w = 16
        for c in range(1, total_excel_cols + 1):
            ws.column_dimensions[get_column_letter(c)].width = col_w

        from app.algorithm.labels import get_seat_label

        # Seat grid column headers
        for c in range(room.columns):
            for si in range(max_spb):
                col_idx = c * max_spb + si + 1
                pos_char = chr(ord('A') + si)
                label = f"Bench {c + 1} ({pos_char})"
                hdr = ws.cell(2, col_idx, label)
                hdr.font = Font(bold=True, size=9)
                hdr.alignment = Alignment(horizontal="center")
                hdr.fill = _make_fill("E2E8F0")

        grid: dict[tuple, SeatAssignment] = {
            (s.row, s.col, s.seat_index): s for s in room_seats
        }

        for r in range(room.rows):
            excel_row = r + 3
            ws.row_dimensions[excel_row].height = 38
            for c in range(room.columns):
                bench_seats = room.get_bench_seats(r, c)
                for si in range(max_spb):
                    col_idx = c * max_spb + si + 1
                    cell = ws.cell(excel_row, col_idx)

                    if si < bench_seats:
                        seat = grid.get((r, c, si))
                        lbl = (seat.seat_label if seat and seat.seat_label else get_seat_label(r, c, si, room.columns))

                        if seat and not seat.is_empty:
                            d_code = seat.department.code if seat.department else (f"D{seat.department_id}" if seat.department_id else "")
                            med = f" ({seat.medium})" if getattr(seat, "medium", None) else ""
                            arr_badge = " [ARR]" if seat.is_arrear else ""
                            cell.value = f"{lbl}\n{seat.roll_number or ''}{arr_badge}\n{d_code}{med}"
                            cell.fill = _dept_fill(seat.department_id)
                        else:
                            cell.value = f"{lbl}\n(Unused)"
                            cell.fill = _make_fill(_EMPTY_FILL)
                        cell.border = _border
                    else:
                        cell.value = "—"
                        cell.fill = _make_fill("FFFFFF")

                    cell.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                    cell.font = Font(size=8)

        # Row labels in helper column
        for r in range(room.rows):
            helper = ws.cell(r + 3, total_excel_cols + 1, f"Row {r + 1}")
            helper.font = Font(size=8, color="888888")
            helper.alignment = Alignment(horizontal="left", vertical="center")

        curr_row = room.rows + 4
        # Student list table
        seated_in_hall = [s for s in room_seats if not s.is_empty]
        if seated_in_hall:
            ws.cell(curr_row, 1, f"Student Desk Allocation List - Hall {room.name}").font = Font(bold=True, size=10)
            curr_row += 1
            tbl_headers = ["Seat Label", "Register Number", "Department", "Medium", "Subject Code", "Candidate Type"]
            for ci, th in enumerate(tbl_headers, 1):
                tc = ws.cell(curr_row, ci, th)
                tc.font = Font(bold=True, color="FFFFFF", size=9)
                tc.fill = _make_fill("334155")
                tc.alignment = Alignment(horizontal="center")

            curr_row += 1
            for s in seated_in_hall:
                lbl = s.seat_label or get_seat_label(s.row, s.col, s.seat_index, room.columns)
                d_code = s.department.code if s.department else (f"Dept {s.department_id}" if s.department_id else "—")
                s_med = getattr(s, "medium", "EN") or "EN"
                s_type = "Arrear" if s.is_arrear else "Regular"
                subj = s.subject_code or s.subject or "—"
                ws.cell(curr_row, 1, lbl).alignment = Alignment(horizontal="center")
                ws.cell(curr_row, 2, s.roll_number or "—").alignment = Alignment(horizontal="center")
                ws.cell(curr_row, 3, d_code).alignment = Alignment(horizontal="center")
                ws.cell(curr_row, 4, s_med).alignment = Alignment(horizontal="center")
                ws.cell(curr_row, 5, subj).alignment = Alignment(horizontal="left")
                ws.cell(curr_row, 6, s_type).alignment = Alignment(horizontal="center")
                curr_row += 1

        curr_row += 1
        sig = ws.cell(curr_row, 1,
                      f"Invigilator 1 Signature: ________________________  "
                      f"Invigilator 2 Signature: ________________________  "
                      f"Date: _______________  Hall: {room.name}")
        sig.font = Font(size=9, italic=True, color="444444")
        sig.alignment = Alignment(horizontal="left")

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws_sum = wb.create_sheet(title="Summary", index=0)
    ws_sum.column_dimensions["A"].width = 20
    ws_sum.column_dimensions["B"].width = 16
    ws_sum.column_dimensions["C"].width = 16
    ws_sum.column_dimensions["D"].width = 16

    headers = ["Room", "Capacity", "Seated", "Utilisation"]
    for ci, h in enumerate(headers, 1):
        c = ws_sum.cell(1, ci, h)
        c.font = Font(bold=True, color=_HEADER_FONT_COLOR)
        c.fill = _make_fill(_HEADER_FILL)
        c.alignment = Alignment(horizontal="center")

    for ri, room in enumerate(rooms, 2):
        room_seats = seats_by_room.get(room.id, [])
        used = sum(1 for s in room_seats if not s.is_empty)
        cap = room.capacity
        ws_sum.cell(ri, 1, room.name)
        ws_sum.cell(ri, 2, cap)
        ws_sum.cell(ri, 3, used)
        ws_sum.cell(ri, 4, f"{round(used / cap * 100, 1)}%" if cap else "—")

    # Warnings
    warn_start = len(rooms) + 3
    ws_sum.cell(warn_start, 1, "Warnings").font = Font(bold=True)
    for i, w in enumerate(plan.warnings, warn_start + 1):
        cell = ws_sum.cell(i, 1, w)
        ws_sum.merge_cells(
            start_row=i, start_column=1, end_row=i, end_column=4
        )
        cell.alignment = Alignment(wrap_text=True)
        ws_sum.row_dimensions[i].height = 30
        if w.startswith("ERROR"):
            cell.font = Font(color="DC2626", size=9)
        elif w.startswith("WARNING"):
            cell.font = Font(color="D97706", size=9)
        else:
            cell.font = Font(color="6B7280", size=9)

    # ── Branch Summary sheet ──────────────────────────────────────────────────
    ws_branch = wb.create_sheet(title="Branch Summary", index=1)
    ws_branch.column_dimensions["A"].width = 16
    ws_branch.column_dimensions["B"].width = 28
    ws_branch.column_dimensions["C"].width = 18
    ws_branch.column_dimensions["D"].width = 24
    ws_branch.column_dimensions["E"].width = 16

    b_headers = ["Department", "Subject", "Hall", "Roll Range", "Candidate Count"]
    for ci, h in enumerate(b_headers, 1):
        c = ws_branch.cell(1, ci, h)
        c.font = Font(bold=True, color=_HEADER_FONT_COLOR)
        c.fill = _make_fill(_HEADER_FILL)
        c.alignment = Alignment(horizontal="center")

    b_row = 2
    for r in rooms:
        if not r:
            continue
        r_seats = seats_by_room.get(r.id, [])
        dept_ids = sorted(list({s.department_id for s in r_seats if not s.is_empty and s.department_id}))
        for d_id in dept_ids:
            dept_seats = [s for s in r_seats if s.department_id == d_id and not s.is_empty]
            dept_seats.sort(key=lambda s: (s.row, s.col, s.seat_index))
            d_code = dept_seats[0].department.code if (dept_seats[0].department) else f"Dept {d_id}"
            subj = dept_seats[0].subject or "Exam"
            rolls = [s.roll_number for s in dept_seats if s.roll_number]
            r_range = f"{rolls[0]} - {rolls[-1]}" if len(rolls) > 1 else (rolls[0] if rolls else "-")

            ws_branch.cell(b_row, 1, d_code).alignment = Alignment(horizontal="center")
            ws_branch.cell(b_row, 2, subj)
            ws_branch.cell(b_row, 3, r.name)
            ws_branch.cell(b_row, 4, r_range).alignment = Alignment(horizontal="center")
            ws_branch.cell(b_row, 5, len(rolls)).alignment = Alignment(horizontal="center")
            b_row += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
