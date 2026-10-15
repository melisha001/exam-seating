"""
routers/plans.py — FastAPI routes for seating plan generation, retrieval,
                   seat swapping, and PDF/Excel export.

Endpoints
---------
POST   /plans/generate               Generate (or regenerate) a plan for a session
GET    /plans/session/{session_id}   Retrieve the plan for a session
PUT    /plans/{plan_id}/swap         Manually swap two seats (re-validates constraints)
GET    /plans/{plan_id}/export/pdf   Download PDF
GET    /plans/{plan_id}/export/excel Download Excel
"""

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.room import Room
from app.models.department import Department
from app.models.seating_plan import SeatingPlan, SeatAssignment

import app.crud.exam_sessions as session_crud
import app.crud.exam_entries as entry_crud
import app.crud.rooms as room_crud
import app.crud.seating_plans as plan_crud

import app.crud.departments as dept_crud

from app.algorithm.seating import generate_seating
from app.algorithm.roll_numbers import generate_rolls
from app.algorithm.models import DeptGroup, RoomSpec
from app.algorithm.adjacency import get_neighbors

from app.services.export_pdf import generate_pdf
from app.services.export_excel import generate_excel

from app.schemas.plan_request import (
    GeneratePlanRequest,
    SwapRequest,
    PlanSummary,
    PlanDetail,
    RoomGrid,
    PlanBranchSummaryResponse,
    BranchSummaryItem,
    BranchHallAllocation,
    HallSummaryItem,
    HallBranchItem,
    MatrixCell,
)

router = APIRouter(prefix="/plans", tags=["Seating Plans"])


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _plan_or_404(db: Session, plan_id: int) -> SeatingPlan:
    plan = plan_crud.get_plan_by_id(db, plan_id)
    if not plan:
        raise HTTPException(status_code=404, detail=f"Plan {plan_id} not found.")
    return plan


def _seat_or_404(db: Session, seat_id: int) -> SeatAssignment:
    seat = plan_crud.get_seat(db, seat_id)
    if not seat:
        raise HTTPException(status_code=404, detail=f"Seat {seat_id} not found.")
    return seat


def _build_plan_detail(plan: SeatingPlan, db: Session) -> PlanDetail:
    """Construct a PlanDetail response from a SeatingPlan ORM object."""
    all_seats = plan_crud.get_seats_for_plan(db, plan.id)

    # Group seats by room
    room_ids_ordered: list[int] = []
    seats_by_room: dict[int, list[SeatAssignment]] = {}
    for seat in all_seats:
        if seat.room_id not in seats_by_room:
            room_ids_ordered.append(seat.room_id)
            seats_by_room[seat.room_id] = []
        seats_by_room[seat.room_id].append(seat)

    room_grids: list[RoomGrid] = []
    for rid in room_ids_ordered:
        room: Room = room_crud.get_room(db, rid)
        r_seats = seats_by_room[rid]
        used = sum(1 for s in r_seats if not s.is_empty)
        cap = room.capacity
        util = round(used / cap * 100, 1) if cap else 0.0

        seat_dicts = [
            {
                "id": s.id,
                "plan_id": s.plan_id,
                "room_id": s.room_id,
                "row": s.row,
                "col": s.col,
                "seat_index": s.seat_index,
                "roll_number": s.roll_number,
                "department_id": s.department_id,
                "subject": s.subject,
                "subject_code": s.subject_code,
                "medium": getattr(s, "medium", "EN") or "EN",
                "is_empty": s.is_empty,
                "is_arrear": s.is_arrear,
                "is_locked": s.is_locked,
                "seat_label": s.seat_label,
            }
            for s in r_seats
        ]
        room_grids.append(RoomGrid(
            room_id=rid,
            room_name=room.name,
            invigilator_name=room.invigilator_name,
            invigilator2_name=room.invigilator2_name,
            invigilator1_subject=room.invigilator1_subject,
            invigilator2_subject=room.invigilator2_subject,
            capacity=cap,
            assigned=used,
            utilization_pct=util,
            rows=room.rows,
            columns=room.columns,
            bench_config=room.bench_config,
            seater_mode=getattr(room, "seater_mode", "auto") or "auto",
            seater_breakdown=room.seater_breakdown,
            seats=seat_dicts,
        ))

    return PlanDetail(
        plan_id=plan.id,
        session_id=plan.session_id,
        rooms_used=len(room_ids_ordered),
        total_assigned=sum(rg.assigned for rg in room_grids),
        unseated_count=plan.unseated_count,
        warnings=plan.warnings,
        rooms=room_grids,
    )


def _validate_swap_constraints(
    seat_a: SeatAssignment,
    seat_b: SeatAssignment,
    all_seats: list[SeatAssignment],
    db: Session,
) -> list[str]:
    """
    Simulate swapping seats A and B and return a list of constraint violations.
    An empty list means the swap is valid.
    """
    # Build a grid of the current plan with the swap applied
    grid: dict[tuple, SeatAssignment] = {}
    for seat in all_seats:
        if seat.id == seat_a.id:
            # Pretend seat_b's student sits here
            mock = SeatAssignment(
                id=seat.id, plan_id=seat.plan_id, room_id=seat.room_id,
                row=seat.row, col=seat.col, seat_index=seat.seat_index,
                roll_number=seat_b.roll_number,
                department_id=seat_b.department_id,
                subject=seat_b.subject,
                is_empty=seat_b.is_empty,
            )
            grid[(seat.room_id, seat.row, seat.col, seat.seat_index)] = mock
        elif seat.id == seat_b.id:
            mock = SeatAssignment(
                id=seat.id, plan_id=seat.plan_id, room_id=seat.room_id,
                row=seat.row, col=seat.col, seat_index=seat.seat_index,
                roll_number=seat_a.roll_number,
                department_id=seat_a.department_id,
                subject=seat_a.subject,
                is_empty=seat_a.is_empty,
            )
            grid[(seat.room_id, seat.row, seat.col, seat.seat_index)] = mock
        else:
            grid[(seat.room_id, seat.row, seat.col, seat.seat_index)] = seat

    violations: list[str] = []

    def check_seat(seat: SeatAssignment, room: Room) -> None:
        if seat.is_empty:
            return
        neighbors = get_neighbors(
            seat.row, seat.col, seat.seat_index,
            room.rows, room.columns, room.seats_per_bench,
            check_front_back=True,
        )
        for nr, nc, ns in neighbors:
            n = grid.get((seat.room_id, nr, nc, ns))
            if n and not n.is_empty:
                if n.department_id == seat.department_id:
                    violations.append(
                        f"Same department adjacent: "
                        f"{seat.roll_number} ↔ {n.roll_number}"
                    )
                elif n.subject and seat.subject and n.subject == seat.subject:
                    violations.append(
                        f"Same subject '{seat.subject}' adjacent: "
                        f"{seat.roll_number} ↔ {n.roll_number}"
                    )

    for target_seat_id in (seat_a.id, seat_b.id):
        target = grid.get(
            (seat_a.room_id, seat_a.row, seat_a.col, seat_a.seat_index)
            if target_seat_id == seat_a.id
            else (seat_b.room_id, seat_b.row, seat_b.col, seat_b.seat_index)
        )
        if target:
            room = room_crud.get_room(db, target.room_id)
            if room:
                check_seat(target, room)

    # Deduplicate (each violation appears once per direction)
    return list(dict.fromkeys(violations))


# ─────────────────────────────────────────────────────────────────────────────
# Routes
# ─────────────────────────────────────────────────────────────────────────────

@router.post("/generate", response_model=PlanDetail, status_code=status.HTTP_201_CREATED)
def generate_plan(body: GeneratePlanRequest, db: Session = Depends(get_db)):
    """
    Generate (or regenerate) a seating plan for a session.

    - Fetches all ExamEntry records for the session.
    - Builds DeptGroup / RoomSpec objects.
    - Runs the greedy algorithm.
    - Persists the result to the database (replaces any existing plan).
    """
    # Validate session
    session = session_crud.get_session(db, body.session_id)
    if not session:
        raise HTTPException(404, f"Session {body.session_id} not found.")

    # Validate and load rooms in requested order
    rooms: list[Room] = []
    for rid in body.room_ids:
        room = room_crud.get_room(db, rid)
        if not room:
            raise HTTPException(422, f"Room {rid} not found.")
        rooms.append(room)

    # Load exam entries
    entries = entry_crud.get_entries(db, session_id=body.session_id)
    if not entries:
        raise HTTPException(
            422,
            f"Session {body.session_id} has no exam entries. "
            "Add departments and student counts before generating a plan.",
        )

    # Build algorithm inputs
    dept_groups: list[DeptGroup] = []
    for entry in entries:
        exclusions_list = (
            [x.strip() for x in entry.exclusions.split(",") if x.strip()]
            if entry.exclusions
            else None
        )
        rolls = generate_rolls(
            entry.department.code,
            entry.student_count,
            entry.roll_start,
            entry.roll_end,
            exclusions=exclusions_list,
        )
        subj_label = f"{entry.subject_code} - {entry.subject}" if entry.subject_code else entry.subject
        dept_groups.append(DeptGroup(
            dept_id=entry.department_id,
            dept_code=entry.department.code,
            subject=subj_label,
            rolls=rolls,
            subject_code=entry.subject_code,
            is_arrear=entry.is_arrear,
            medium=getattr(entry, "medium", "EN") or "EN",
        ))

    room_specs: list[RoomSpec] = [
        RoomSpec(
            room_id=r.id,
            room_name=r.name,
            rows=r.rows,
            cols=r.columns,
            seats_per_bench=r.seats_per_bench,
            invigilator_name=r.invigilator_name,
            invigilator2_name=r.invigilator2_name,
            invigilator1_subject=r.invigilator1_subject,
            invigilator2_subject=r.invigilator2_subject,
            bench_config=r.bench_config,
            seater_mode=getattr(r, "seater_mode", "auto") or "auto",
        )
        for r in rooms
    ]

    # Run algorithm
    try:
        result = generate_seating(
            dept_groups,
            room_specs,
            check_front_back=body.check_front_back,
            treat_different_medium_as_different_paper=getattr(body, "treat_different_medium_as_different_paper", False),
            equal_students_per_hall=getattr(body, "equal_students_per_hall", False),
        )
    except ValueError as ve:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=str(ve),
        )

    # Persist
    plan = plan_crud.save_plan(db, body.session_id, result)
    return _build_plan_detail(plan, db)


class TargetCheckRequest(BaseModel):
    session_id: int
    room_ids: list[int]
    equal_students_per_hall: bool = False


@router.post("/target-check")
def check_targets(body: TargetCheckRequest, db: Session = Depends(get_db)):
    """Check target student counts vs current capacities for selected halls."""
    from app.algorithm.seating import get_hall_targets

    session = session_crud.get_session(db, body.session_id)
    if not session:
        raise HTTPException(404, f"Session {body.session_id} not found.")

    entries = entry_crud.get_entries(db, session_id=body.session_id)
    total_students = sum(e.student_count for e in entries)
    if not body.room_ids:
        return {
            "is_valid": False,
            "total_students": total_students,
            "total_capacity": 0,
            "shortage": total_students,
            "messages": ["No examination halls selected."],
            "halls": [],
        }

    rooms = [room_crud.get_room(db, rid) for rid in body.room_ids]
    valid_rooms = [r for r in rooms if r is not None]
    total_capacity = sum(r.capacity for r in valid_rooms)

    room_specs = [
        RoomSpec(
            room_id=r.id,
            room_name=r.name,
            rows=r.rows,
            cols=r.columns,
            seats_per_bench=r.seats_per_bench,
            bench_config=r.bench_config,
            seater_mode=getattr(r, "seater_mode", "auto") or "auto",
        )
        for r in valid_rooms
    ]

    if total_students > total_capacity:
        shortage = total_students - total_capacity
        return {
            "is_valid": False,
            "total_students": total_students,
            "total_capacity": total_capacity,
            "shortage": shortage,
            "messages": [
                f"Short by {shortage} seat{'s' if shortage > 1 else ''}. Add a hall or edit benches in Halls & Depts."
            ],
            "halls": [],
        }

    targets = get_hall_targets(
        total_students,
        len(valid_rooms),
        rooms=room_specs,
        equal_students_per_hall=body.equal_students_per_hall,
    )
    halls_info = []
    messages = []
    for r, target in zip(valid_rooms, targets):
        cap = r.capacity
        diff = cap - target
        is_matched = (cap == target)
        halls_info.append({
            "room_id": r.id,
            "room_name": r.name,
            "capacity": cap,
            "target": target,
            "diff": diff,
            "is_matched": is_matched,
            "seater_breakdown": r.seater_breakdown,
        })

    if total_students < total_capacity:
        unused = total_capacity - total_students
        messages.append(f"{unused} seats will stay unused. Edit halls in Halls & Depts to match, or continue.")

    return {
        "is_valid": True,
        "total_students": total_students,
        "total_capacity": total_capacity,
        "shortage": 0,
        "messages": messages,
        "halls": halls_info,
    }



@router.get("/session/{session_id}", response_model=PlanDetail)
def get_plan_by_session(session_id: int, db: Session = Depends(get_db)):
    """Retrieve the seating plan for a given session."""
    plan = plan_crud.get_plan_by_session(db, session_id)
    if not plan:
        raise HTTPException(
            404,
            f"No seating plan found for session {session_id}. "
            "Generate one first via POST /plans/generate.",
        )
    return _build_plan_detail(plan, db)


@router.put("/{plan_id}/swap", response_model=dict)
def swap_seats(plan_id: int, body: SwapRequest, db: Session = Depends(get_db)):
    """
    Manually swap the students in two seats.

    Validates all adjacency constraints after the swap.
    Returns 409 with violation details if the swap would break a rule.
    """
    plan = _plan_or_404(db, plan_id)
    seat_a = _seat_or_404(db, body.seat_id_a)
    seat_b = _seat_or_404(db, body.seat_id_b)

    # Both seats must belong to this plan
    for seat in (seat_a, seat_b):
        if seat.plan_id != plan_id:
            raise HTTPException(
                422,
                f"Seat {seat.id} does not belong to plan {plan_id}.",
            )

    all_seats = plan_crud.get_seats_for_plan(db, plan_id)
    violations = _validate_swap_constraints(seat_a, seat_b, all_seats, db)

    if violations:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={
                "message": "Swap rejected: constraint violation(s) detected.",
                "violations": violations,
            },
        )

    # Apply the swap
    (
        seat_a.roll_number, seat_b.roll_number,
        seat_a.department_id, seat_b.department_id,
        seat_a.subject, seat_b.subject,
        seat_a.subject_code, seat_b.subject_code,
        seat_a.is_empty, seat_b.is_empty,
        seat_a.is_arrear, seat_b.is_arrear,
    ) = (
        seat_b.roll_number, seat_a.roll_number,
        seat_b.department_id, seat_a.department_id,
        seat_b.subject, seat_a.subject,
        seat_b.subject_code, seat_a.subject_code,
        seat_b.is_empty, seat_a.is_empty,
        seat_b.is_arrear, seat_a.is_arrear,
    )
    db.commit()

    return {
        "message": "Swap applied successfully.",
        "seat_a_id": seat_a.id,
        "seat_b_id": seat_b.id,
    }


# ── Export endpoints ──────────────────────────────────────────────────────────

@router.get("/{plan_id}/export/pdf")
def export_pdf(plan_id: int, room_id: int | None = None, db: Session = Depends(get_db)):
    """Download the seating plan as a PDF (one landscape A4 page per room, or filtered to a single room)."""
    plan = _plan_or_404(db, plan_id)
    all_seats = plan_crud.get_seats_for_plan(db, plan_id)

    # Collect unique room IDs in order
    room_ids_ordered: list[int] = []
    seats_by_room: dict[int, list[SeatAssignment]] = {}
    for seat in all_seats:
        if room_id is not None and seat.room_id != room_id:
            continue
        if seat.room_id not in seats_by_room:
            room_ids_ordered.append(seat.room_id)
            seats_by_room[seat.room_id] = []
        seats_by_room[seat.room_id].append(seat)

    rooms = [room_crud.get_room(db, rid) for rid in room_ids_ordered]
    rooms = [r for r in rooms if r is not None]

    pdf_bytes = generate_pdf(plan, rooms, seats_by_room)
    filename = f"seating_plan_{plan_id}_room_{room_id}.pdf" if room_id else f"seating_plan_{plan_id}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


@router.get("/{plan_id}/export/excel")
def export_excel(plan_id: int, room_id: int | None = None, db: Session = Depends(get_db)):
    """Download the seating plan as an Excel workbook (or single room)."""
    plan = _plan_or_404(db, plan_id)
    all_seats = plan_crud.get_seats_for_plan(db, plan_id)

    room_ids_ordered: list[int] = []
    seats_by_room: dict[int, list[SeatAssignment]] = {}
    for seat in all_seats:
        if room_id is not None and seat.room_id != room_id:
            continue
        if seat.room_id not in seats_by_room:
            room_ids_ordered.append(seat.room_id)
            seats_by_room[seat.room_id] = []
        seats_by_room[seat.room_id].append(seat)

    rooms = [room_crud.get_room(db, rid) for rid in room_ids_ordered]
    rooms = [r for r in rooms if r is not None]

    xlsx_bytes = generate_excel(plan, rooms, seats_by_room)
    filename = f"seating_plan_{plan_id}_room_{room_id}.xlsx" if room_id else f"seating_plan_{plan_id}.xlsx"
    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


# ── Branch-wise Summary & Session Plan Cleanup ────────────────────────────────

def _build_branch_summary(plan: SeatingPlan, db: Session) -> PlanBranchSummaryResponse:
    all_seats = plan_crud.get_seats_for_plan(db, plan.id)
    departments = dept_crud.get_departments(db)
    dept_map = {d.id: d for d in departments}

    room_ids_ordered: list[int] = []
    seats_by_room: dict[int, list[SeatAssignment]] = {}
    for seat in all_seats:
        if seat.room_id not in seats_by_room:
            room_ids_ordered.append(seat.room_id)
            seats_by_room[seat.room_id] = []
        seats_by_room[seat.room_id].append(seat)

    room_map = {rid: room_crud.get_room(db, rid) for rid in room_ids_ordered}

    # Group seated non-empty seats by (dept_id, room_id)
    dept_room_seats: dict[int, dict[int, list[SeatAssignment]]] = {}
    for seat in all_seats:
        if not seat.is_empty and seat.department_id:
            dept_room_seats.setdefault(seat.department_id, {}).setdefault(seat.room_id, []).append(seat)

    entries = entry_crud.get_entries(db, session_id=plan.session_id)
    dept_subject_map = {
        e.department_id: (f"{e.subject_code} - {e.subject}" if e.subject_code else e.subject)
        for e in entries
    }

    branch_items: list[BranchSummaryItem] = []
    matrix_data: dict[str, dict[str, MatrixCell]] = {}
    department_codes: list[str] = []

    for dept in departments:
        if dept.id not in dept_room_seats:
            continue
        department_codes.append(dept.code)
        matrix_data[dept.code] = {}

        total_dept_students = 0
        halls_alloc: list[BranchHallAllocation] = []
        subj = dept_subject_map.get(dept.id, "Exam")

        for rid in room_ids_ordered:
            rm = room_map.get(rid)
            rm_name = rm.name if rm else f"Room {rid}"
            st_list = dept_room_seats[dept.id].get(rid, [])

            st_list.sort(key=lambda s: (s.row, s.col, s.seat_index))
            rolls = [s.roll_number for s in st_list if s.roll_number]
            count = len(rolls)
            total_dept_students += count

            if count > 1:
                r_range = f"{rolls[0]} - {rolls[-1]}"
            elif count == 1:
                r_range = rolls[0]
            else:
                r_range = "-"

            matrix_data[dept.code][rm_name] = MatrixCell(student_count=count, roll_range=r_range)

            if count > 0:
                halls_alloc.append(BranchHallAllocation(
                    room_id=rid,
                    room_name=rm_name,
                    invigilator_name=rm.invigilator_name if rm else None,
                    student_count=count,
                    roll_range=r_range,
                    roll_numbers=rolls,
                    seats=[f"R{s.row+1}C{s.col+1}S{s.seat_index+1}" for s in st_list],
                ))

        branch_items.append(BranchSummaryItem(
            department_id=dept.id,
            department_code=dept.code,
            department_name=dept.name,
            subject=subj,
            total_students=total_dept_students,
            halls=halls_alloc,
        ))

    hall_items: list[HallSummaryItem] = []
    room_names: list[str] = []

    for rid in room_ids_ordered:
        rm = room_map.get(rid)
        if not rm:
            continue
        room_names.append(rm.name)
        rm_seats = seats_by_room.get(rid, [])
        assigned = sum(1 for s in rm_seats if not s.is_empty)
        cap = rm.capacity
        util = round(assigned / cap * 100, 1) if cap > 0 else 0.0

        b_items: list[HallBranchItem] = []
        for dept in departments:
            st_list = dept_room_seats.get(dept.id, {}).get(rid, [])
            if st_list:
                st_list.sort(key=lambda s: (s.row, s.col, s.seat_index))
                rolls = [s.roll_number for s in st_list if s.roll_number]
                r_range = f"{rolls[0]} - {rolls[-1]}" if len(rolls) > 1 else (rolls[0] if rolls else "-")
                b_items.append(HallBranchItem(
                    department_id=dept.id,
                    department_code=dept.code,
                    department_name=dept.name,
                    student_count=len(rolls),
                    roll_range=r_range,
                ))

        hall_items.append(HallSummaryItem(
            room_id=rid,
            room_name=rm.name,
            invigilator_name=rm.invigilator_name,
            capacity=cap,
            total_seated=assigned,
            utilization_pct=util,
            branches=b_items,
        ))

    total_assigned = sum(b.total_students for b in branch_items)

    return PlanBranchSummaryResponse(
        plan_id=plan.id,
        session_id=plan.session_id,
        total_assigned=total_assigned,
        unseated_count=plan.unseated_count,
        warnings=plan.warnings,
        departments=branch_items,
        halls=hall_items,
        department_codes=department_codes,
        room_names=room_names,
        matrix=matrix_data,
    )


@router.get("/{plan_id}/summary", response_model=PlanBranchSummaryResponse)
def get_plan_branch_summary(plan_id: int, db: Session = Depends(get_db)):
    """Retrieve branch-wise and hall-wise allocation summary for a plan."""
    plan = _plan_or_404(db, plan_id)
    return _build_branch_summary(plan, db)


@router.get("/{plan_id}/export/overall-pdf")
def export_overall_pdf(plan_id: int, db: Session = Depends(get_db)):
    """
    Download overall seating arrangement PDF with exactly 3 columns:
    Department Name, Register Number Range (Start - End), Hall Name.
    """
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, portrait
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.units import cm, mm

    plan = _plan_or_404(db, plan_id)
    all_seats = plan_crud.get_seats_for_plan(db, plan_id)
    departments = dept_crud.get_departments(db)
    dept_map = {d.id: d for d in departments}

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=portrait(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title="Overall Seating Arrangement",
    )
    styles = getSampleStyleSheet()
    heading = ParagraphStyle("heading", parent=styles["Heading1"], fontSize=14, spaceAfter=6, alignment=1)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=9, textColor=colors.grey, alignment=1)

    story = [
        Paragraph("OVERALL EXAMINATION SEATING ARRANGEMENT", heading),
        Paragraph(f"Plan #{plan.id} | Session ID: {plan.session_id}", sub),
        Spacer(1, 6 * mm),
    ]

    # Exactly 3 columns
    table_rows = [["Department Name", "Register Number Range", "Hall Name"]]

    # Collect grouped allocations: by department, then hall
    dept_room_seats: dict[int, dict[int, list[SeatAssignment]]] = {}
    for seat in all_seats:
        if not seat.is_empty and seat.department_id:
            dept_room_seats.setdefault(seat.department_id, {}).setdefault(seat.room_id, []).append(seat)

    for dept in departments:
        rooms_for_dept = dept_room_seats.get(dept.id, {})
        for rid, s_list in rooms_for_dept.items():
            room = room_crud.get_room(db, rid)
            room_name = room.name if room else f"Hall {rid}"
            s_list.sort(key=lambda s: (s.row, s.col, s.seat_index))
            rolls = [s.roll_number for s in s_list if s.roll_number]
            if rolls:
                roll_range = f"{rolls[0]} - {rolls[-1]}" if len(rolls) > 1 else rolls[0]
                table_rows.append([dept.name, roll_range, room_name])

    if len(table_rows) == 1:
        table_rows.append(["—", "No students allocated", "—"])

    tbl = Table(table_rows, colWidths=[6 * cm, 7 * cm, 5 * cm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(tbl)
    doc.build(story)

    return Response(
        content=buf.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="overall_seating_arrangement_{plan_id}.pdf"'},
    )


@router.get("/{plan_id}/export/overall-excel")
def export_overall_excel(plan_id: int, db: Session = Depends(get_db)):
    """
    Download overall seating arrangement Excel with exactly 3 columns:
    Department Name, Register Number Range, Hall Name.
    """
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    plan = _plan_or_404(db, plan_id)
    all_seats = plan_crud.get_seats_for_plan(db, plan_id)
    departments = dept_crud.get_departments(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Overall Arrangement"

    headers = ["Department Name", "Register Number Range", "Hall Name"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 25
    ws.row_dimensions[1].height = 24

    dept_room_seats: dict[int, dict[int, list[SeatAssignment]]] = {}
    for seat in all_seats:
        if not seat.is_empty and seat.department_id:
            dept_room_seats.setdefault(seat.department_id, {}).setdefault(seat.room_id, []).append(seat)

    row_idx = 2
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for dept in departments:
        rooms_for_dept = dept_room_seats.get(dept.id, {})
        for rid, s_list in rooms_for_dept.items():
            room = room_crud.get_room(db, rid)
            room_name = room.name if room else f"Hall {rid}"
            s_list.sort(key=lambda s: (s.row, s.col, s.seat_index))
            rolls = [s.roll_number for s in s_list if s.roll_number]
            if rolls:
                roll_range = f"{rolls[0]} - {rolls[-1]}" if len(rolls) > 1 else rolls[0]
                ws.cell(row_idx, 1, dept.name).border = border
                ws.cell(row_idx, 2, roll_range).border = border
                ws.cell(row_idx, 3, room_name).border = border
                ws.cell(row_idx, 1).alignment = Alignment(horizontal="left", vertical="center")
                ws.cell(row_idx, 2).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row_idx, 3).alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[row_idx].height = 20
                row_idx += 1

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="overall_seating_arrangement_{plan_id}.xlsx"'},
    )


@router.post("/{plan_id}/save-changes", response_model=dict)
def save_plan_changes(plan_id: int, body: dict = None, db: Session = Depends(get_db)):
    """Persist and lock seating arrangement changes."""
    plan = _plan_or_404(db, plan_id)
    all_seats = plan_crud.get_seats_for_plan(db, plan_id)
    for seat in all_seats:
        seat.is_locked = True
    db.commit()
    return {"message": "All seating assignments saved and locked successfully.", "plan_id": plan_id}


@router.delete("/session/{session_id}", response_model=dict)
def clear_plan_for_session(session_id: int, db: Session = Depends(get_db)):
    """Clear seating plan only for a session, keeping session and exam entries intact."""
    session = session_crud.get_session(db, session_id)
    if not session:
        raise HTTPException(404, f"Session {session_id} not found.")

    plan = plan_crud.get_plan_by_session(db, session_id)
    if not plan:
        return {"message": f"No seating plan exists for session {session_id}."}

    plan_crud.delete_plan(db, plan)
    return {"message": f"Seating plan for session {session_id} cleared successfully."}

